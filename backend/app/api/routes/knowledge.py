from fastapi import APIRouter, HTTPException, Request
from typing import Optional, List, Dict, Any
from pathlib import Path

from app.rag.ingestion import ingestion_service
from app.rag.retrieval.service import retrieval_service
from app.rag.retrieval.models import SearchCompanyPolicyRequest
from app.core.logging import get_logger

logger = get_logger(__name__)

router = APIRouter(tags=["knowledge"])

DOCS_DIR = Path(__file__).resolve().parent.parent.parent.parent.parent / "data" / "documents"

SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx"}


class UnsupportedFileType(ValueError):
    """Raised when the uploaded filename extension is not supported."""


def _extract_pdf_text(data: bytes) -> str:
    import io
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                pass
            if reader.is_encrypted:
                raise ValueError("Password-protected PDF. Remove the password and upload again.")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ValueError:
        raise
    except PdfReadError as e:
        raise ValueError(f"Corrupt or unreadable PDF: {e}")
    except Exception as e:
        raise ValueError(f"Failed to read PDF: {e}")


def _extract_docx_text(data: bytes) -> str:
    import io
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ValueError(f"Corrupt or unreadable DOCX: {e}")

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _extract_text(file_content: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        content = _extract_pdf_text(file_content)
    elif ext == ".docx":
        content = _extract_docx_text(file_content)
    elif ext in (".md", ".txt"):
        try:
            content = file_content.decode("utf-8", errors="replace")
        except Exception as e:
            raise ValueError(f"Failed to decode {ext} file as text: {e}")
    else:
        raise UnsupportedFileType("Unsupported file type. Supported: PDF, DOCX, MD, TXT")

    if not content or not content.strip():
        raise ValueError("No extractable text found (scanned PDF?)")
    return content


@router.get("/documents/available")
async def list_policy_documents():
    """List available policy documents in the data/documents directory."""
    if not DOCS_DIR.exists():
        return {"available": []}
    files = [f.name for f in sorted(DOCS_DIR.glob("*.md"))]
    return {"available": files}


@router.get("/documents")
async def list_ingested_documents():
    """List ingested documents (excluding test-policy) from Supabase."""
    from app.services.database import get_supabase_client
    client = get_supabase_client()
    try:
        result = client.table("documents").select("*").in_("source", ["internal-policy", "external-policy"]).execute()
        documents = result.data or []

        for doc in documents:
            chunk_result = client.table("document_chunks").select("id", count="exact").eq("document_id", doc["id"]).execute()
            doc["chunk_count"] = chunk_result.count if chunk_result.count else 0

        return {"data": documents, "total": len(documents), "page": 1, "page_size": len(documents)}
    except Exception:
        return {"data": [], "total": 0, "page": 1, "page_size": 0}


@router.post("/documents/ingest")
async def ingest_document(request: Request):
    """Ingest a policy document from the data/documents directory or uploaded file."""
    try:
        content_type = request.headers.get("content-type", "")
        if "multipart/form-data" in content_type:
            form = await request.form()
            file = form.get("file")
            filename = form.get("document_name") or (file.filename if file else None)
            source = form.get("source", "internal-policy")
            version = form.get("version", "1.0")
            if not file or not filename:
                raise HTTPException(status_code=400, detail="file and document_name are required")
            file_content = await file.read()
            try:
                content = _extract_text(file_content, filename)
            except UnsupportedFileType as e:
                raise HTTPException(status_code=400, detail=str(e))
            except ValueError as e:
                raise HTTPException(status_code=422, detail=str(e))
            except ImportError as e:
                logger.exception("Missing extraction dependency for %s", filename)
                raise HTTPException(
                    status_code=500,
                    detail=f"Server is missing a required package: {e}. Install it and restart the backend.",
                )
            from app.rag.models import Document, DocumentSource
            import uuid
            doc_id = str(uuid.uuid4())
            document = Document(
                id=doc_id,
                name=filename,
                source=DocumentSource(source),
                version=version,
                content=content,
                metadata={},
            )
            result = await ingestion_service.ingest_document(document)
            return result.model_dump()
        else:
            body = await request.json()
            filename = body.get("filename")
            document_name = body.get("document_name")
            source = body.get("source", "internal-policy")
            version = body.get("version", "1.0")
            if not filename:
                raise HTTPException(status_code=400, detail="filename is required")
            file_path = DOCS_DIR / filename
            if not file_path.exists():
                raise HTTPException(status_code=404, detail=f"Document file not found: {filename}")
            result = await ingestion_service.ingest_from_file(
                file_path=str(file_path),
                document_name=document_name or filename,
                source=source,
                version=version,
            )
            return result.model_dump()
    except HTTPException:
        raise
    except Exception:
        logger.exception("Unexpected failure in document ingest")
        raise HTTPException(status_code=500, detail="Internal error while ingesting document")


@router.post("/documents/ingest-all")
async def ingest_all_documents():
    """Ingest all policy documents from the data/documents directory."""
    if not DOCS_DIR.exists():
        raise HTTPException(status_code=404, detail="Documents directory not found")

    results = []
    for md_file in sorted(DOCS_DIR.glob("*.md")):
        try:
            result = await ingestion_service.ingest_from_file(
                file_path=str(md_file),
                document_name=md_file.stem.replace("_", " ").title(),
                source="internal-policy",
                version="1.0",
            )
            results.append({"file": md_file.name, "success": result.success, "chunks": result.chunks_created})
        except Exception as e:
            results.append({"file": md_file.name, "success": False, "error": str(e)})

    return {"results": results}


@router.post("/search")
async def search_policies(body: Dict[str, Any]):
    """Search company policies using RAG retrieval."""
    try:
        query = body.get("query", "")
        max_results = body.get("max_results", 5)

        if not query:
            raise HTTPException(status_code=400, detail="query is required")

        retrieval_request = SearchCompanyPolicyRequest(query=query, max_results=max_results)
        response = await retrieval_service.search_company_policy(retrieval_request)

        return {
            "results": [r.model_dump() for r in response.results],
            "error": response.error,
            "query": query,
            "total_results": len(response.results),
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/documents/{document_id}")
async def get_document(document_id: str):
    from app.services.database import get_supabase_client
    client = get_supabase_client()
    result = client.table("documents").select("*").eq("id", document_id).single().execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Document not found")
    doc = result.data
    chunks_result = client.table("document_chunks").select("*").eq("document_id", document_id).execute()
    doc["chunks"] = chunks_result.data or []
    doc["chunk_count"] = len(chunks_result.data or [])
    return doc
